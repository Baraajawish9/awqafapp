from io import BytesIO
import json
from datetime import datetime, time, timedelta
from unittest.mock import patch

from django.contrib.auth.models import User
from django.http import HttpResponse
from django.core.files.uploadedfile import SimpleUploadedFile
from django.template.loader import render_to_string
from django.test import TestCase
from django.urls import reverse
from django.utils.timezone import get_current_timezone, make_aware
from openpyxl import Workbook, load_workbook

from screen.forms import StudentForm
from screen.models import ExamResult, ScreenSettings, Student
from screen.views import EXAM_TYPE_MAP, parse_room_number


class ExportStudentsExcelTests(TestCase):
    def setUp(self):
        user = User.objects.create_user(username='admin-export', password='12345678')
        self.client.force_login(user)

    def test_export_with_grade_includes_final_grade_and_result(self):
        student = Student.objects.create(
            name='Export Student',
            father_name='Father',
            institute_name='Institute A',
            exam_type='gh',
            memorized_parts='1',
            room='1',
            position=1,
        )
        ExamResult.objects.create(
            number=student.number,
            name=student.name,
            grade=92.5,
            result='ناجح',
            room=1,
            sub_room='0',
        )

        response = self.client.get(reverse('export_students_excel'), {
            'institute': 'Institute A',
            'include_grade': '1',
        })

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        student_row = next(row for row in rows if row and row[1] == 'Export Student')

        self.assertEqual(student_row[8], 92.5)
        self.assertEqual(student_row[9], 'ناجح')

    def test_export_without_grade_keeps_original_columns(self):
        Student.objects.create(
            name='No Grade Export Student',
            institute_name='Institute A',
            exam_type='gh',
            room='1',
            position=1,
        )

        response = self.client.get(reverse('export_students_excel'), {
            'institute': 'Institute A',
            'include_grade': '0',
        })

        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        header_row = next(row for row in sheet.iter_rows(values_only=True) if row and row[0] == 'الرقم')

        self.assertEqual(header_row[:8], ('الرقم', 'الاسم والكنية', 'اسم الأب', 'تاريخ الولادة', 'اسم المعهد', 'غيباً/نظراً', 'الأجزاء المحفوظة', 'الوقت المتوقع'))
        self.assertNotIn('الدرجة', header_row)
        self.assertNotIn('النتيجة', header_row)

    def test_export_uses_one_estimated_time_range_for_group(self):
        ScreenSettings.objects.create(
            room_count=1,
            waiting_count=5,
            estimate_time_per_student=3,
            exam_start_time=time(7, 30),
        )
        for index in range(30):
            Student.objects.create(
                name=f'Range Student {index + 1}',
                institute_name='Institute Range',
                exam_type='gh',
                room='1',
                position=index,
            )

        response = self.client.get(reverse('export_students_excel'), {
            'institute': 'Institute Range',
            'include_grade': '0',
        })

        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        student_rows = [
            row for row in sheet.iter_rows(values_only=True)
            if row and isinstance(row[1], str) and row[1].startswith('Range Student')
        ]

        self.assertEqual(len(student_rows), 30)
        self.assertEqual({row[7] for row in student_rows}, {'07:30 - 09:00'})


class DashboardLateStatusTests(TestCase):
    def test_automatic_status_does_not_promote_top_student_to_in_exam(self):
        User.objects.create_user(username='admin-waiting', password='12345678')
        ScreenSettings.objects.create(room_count=1, waiting_count=5)
        students = [
            Student.objects.create(name='Top Waiting', room='1', status='waiting', position=0),
            Student.objects.create(name='Second Waiting', room='1', status='on_waiting_list', position=1),
        ]

        self.client.login(username='admin-waiting', password='12345678')
        response = self.client.post(reverse('apply_automatic_status'))
        for student in students:
            student.refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertEqual(students[0].status, 'waiting')
        self.assertEqual(students[1].status, 'waiting')

    def test_dashboard_late_status_is_preserved_after_reordering(self):
        User.objects.create_user(username='admin-late', password='12345678')
        students = [
            Student.objects.create(name='Current', room='1', status='in_exam', position=0),
            Student.objects.create(name='Late Student', room='1', status='waiting', position=1),
            Student.objects.create(name='Next Student', room='1', status='waiting', position=2),
        ]

        self.client.login(username='admin-late', password='12345678')
        response = self.client.post(
            reverse('update_student_status', kwargs={'student_number': students[1].number}),
            {'status': 'late'},
        )
        students[1].refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertEqual(students[1].status, 'late')


class ScreenSettingsRoomUserTests(TestCase):
    def test_saving_room_count_creates_branch_logins(self):
        user = User.objects.create_user(username='settings-admin', password='12345678', is_staff=True)
        ScreenSettings.objects.create(room_count=10, waiting_count=5)
        self.client.force_login(user)

        response = self.client.post(reverse('edit_settings'), {
            'room_count': 15,
            'waiting_count': 5,
            'estimate_time_per_student': 5,
            'exam_start_time': '07:00',
            'result_display_seconds': 30,
            'public_screen_mode': 'rooms',
        })

        self.assertRedirects(response, reverse('add_student'), fetch_redirect_response=False)
        self.assertTrue(User.objects.filter(username='room15-1', is_active=True).exists())
        self.assertTrue(User.objects.filter(username='room15-2', is_active=True).exists())

    def test_saving_room_count_does_not_reset_existing_room_passwords(self):
        room_user = User.objects.create_user(username='room1-1', password='custom-pass')
        user = User.objects.create_user(username='settings-admin-password', password='12345678', is_staff=True)
        ScreenSettings.objects.create(room_count=1, waiting_count=5)
        self.client.force_login(user)

        response = self.client.post(reverse('edit_settings'), {
            'room_count': 2,
            'waiting_count': 5,
            'estimate_time_per_student': 5,
            'exam_start_time': '07:00',
            'result_display_seconds': 30,
            'public_screen_mode': 'rooms',
        })

        room_user.refresh_from_db()
        self.assertRedirects(response, reverse('add_student'), fetch_redirect_response=False)
        self.assertTrue(room_user.check_password('custom-pass'))
        self.assertTrue(User.objects.get(username='room2-1').check_password('12345678'))


class PublicScreenModeTests(TestCase):
    def test_rooms_mode_renders_finished_results_page(self):
        student = Student.objects.create(name='Rendered Finished Student', room='1', status='finished')
        result = ExamResult.objects.create(
            number=student.number,
            name=student.name,
            grade=91,
            result='ناجح',
            room=1,
            sub_room='0',
        )
        result.student_name = student.name
        result.result_class = 'success'
        result.display_result = 'ناجح'
        result.display_grade = '91'

        html = render_to_string('screen/public_screen.html', {
            'screen_mode': 'rooms',
            'room_pages': [],
            'room_page_count': 0,
            'screen_page_count': 1,
            'finished_results': [result],
            'finished_results_page_size': 24,
            'rotate_interval_ms': 1200,
            'student_slice_rotate_ms': 4500,
            'visible_students_per_room': 7,
            'tv_density_class': 'tv-density-6 tv-grid-3x2',
            'exam_start_minutes': 420,
            'estimate_minutes': 5,
        })

        self.assertIn('الطلاب المنتهون', html)
        self.assertIn('Rendered Finished Student', html)
        self.assertIn('ناجح', html)
        self.assertIn('إعادة', html)

    def test_rooms_mode_adds_capped_finished_results_page(self):
        user = User.objects.create_user(username='screen-finished-page-admin', password='12345678', is_staff=True)
        ScreenSettings.objects.create(
            room_count=1,
            waiting_count=5,
            estimate_time_per_student=5,
            exam_start_time=time(8, 0),
            public_screen_mode='rooms',
        )

        for index in range(26):
            student = Student.objects.create(
                name=f'Finished Page Student {index + 1}',
                room='1',
                status='finished',
                position=index,
            )
            result = 'ناجح' if index == 25 else 'إعادة'
            ExamResult.objects.create(
                number=student.number,
                name=student.name,
                grade=95 if result == 'ناجح' else 70,
                result=result,
                room=1,
                sub_room='0',
            )

        captured = {}

        def capture_render(request, template_name, context):
            captured.update(context)
            return HttpResponse('ok')

        self.client.force_login(user)
        with patch('screen.views.render', side_effect=capture_render):
            response = self.client.get(reverse('public_screen'))

        finished_results = captured['finished_results']

        self.assertEqual(response.status_code, 200)
        self.assertEqual(captured['screen_page_count'], captured['room_page_count'] + 1)
        self.assertEqual(len(finished_results), 24)
        self.assertEqual(finished_results[0].student_name, 'Finished Page Student 26')
        self.assertEqual(finished_results[0].result_class, 'success')
        self.assertEqual(finished_results[0].display_result, 'ناجح')
        self.assertEqual(finished_results[-1].student_name, 'Finished Page Student 3')
        self.assertNotIn('Finished Page Student 1', [result.student_name for result in finished_results])

    def test_rooms_mode_does_not_add_finished_page_when_empty(self):
        user = User.objects.create_user(username='screen-no-finished-page-admin', password='12345678', is_staff=True)
        ScreenSettings.objects.create(room_count=1, waiting_count=5, public_screen_mode='rooms')
        Student.objects.create(name='Waiting Only Student', room='1', status='waiting')
        captured = {}

        def capture_render(request, template_name, context):
            captured.update(context)
            return HttpResponse('ok')

        self.client.force_login(user)
        with patch('screen.views.render', side_effect=capture_render):
            response = self.client.get(reverse('public_screen'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(captured['finished_results'], [])
        self.assertEqual(captured['screen_page_count'], captured['room_page_count'])

    def test_rooms_mode_keeps_finished_results_off_room_pages(self):
        user = User.objects.create_user(username='screen-results-admin', password='12345678', is_staff=True)
        ScreenSettings.objects.create(
            room_count=1,
            waiting_count=5,
            estimate_time_per_student=5,
            exam_start_time=time(8, 0),
            result_display_seconds=30,
            public_screen_mode='rooms',
        )
        finished_student = Student.objects.create(name='Finished Student', room='1', status='finished', position=0)
        retry_student = Student.objects.create(name='Retry Student', room='1', status='finished', position=1)
        expired_student = Student.objects.create(name='Expired Student', room='1', status='finished', position=2)
        waiting_student = Student.objects.create(name='Waiting Student', room='1', status='waiting', position=2)
        ExamResult.objects.create(
            number=finished_student.number,
            name=finished_student.name,
            grade=100,
            result='ناجح',
            room=1,
            sub_room='0',
        )
        retry_result = ExamResult.objects.create(
            number=retry_student.number,
            name=retry_student.name,
            grade=70,
            result='إعادة',
            room=1,
            sub_room='0',
        )
        expired_result = ExamResult.objects.create(
            number=expired_student.number,
            name=expired_student.name,
            grade=60,
            result='إعادة',
            room=1,
            sub_room='0',
        )
        ExamResult.objects.filter(id=expired_result.id).update(timestamp=make_aware(datetime.now() - timedelta(seconds=31), timezone=get_current_timezone()))
        captured = {}

        def capture_render(request, template_name, context):
            captured.update(context)
            return HttpResponse('ok')

        self.client.force_login(user)
        with patch('screen.views.render', side_effect=capture_render):
            response = self.client.get(reverse('public_screen'))

        room_students = captured['room_pages'][0]['rooms'][0]['students']
        room_card = captured['room_pages'][0]['rooms'][0]
        waiting = next(student for student in room_students if student.number == waiting_student.number)
        room_result_names = [student.name for student in room_students if getattr(student, 'status', '') == 'finished']
        finished_result_names = [result.student_name for result in captured['finished_results']]
        retry = next(result for result in captured['finished_results'] if result.number == retry_student.number)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(room_result_names, [])
        self.assertIn(finished_student.name, finished_result_names)
        self.assertIn(retry_student.name, finished_result_names)
        self.assertIn(expired_student.name, finished_result_names)
        self.assertEqual(room_card['active_count'], 1)
        self.assertEqual(room_card['result_count'], 0)
        self.assertEqual(room_card['total_count'], 1)
        self.assertEqual(retry_result.result, 'إعادة')
        self.assertEqual(retry.result_class, 'retry')
        self.assertEqual(captured['room_page_count'], 1)
        self.assertEqual(captured['screen_page_count'], 2)
        self.assertEqual(waiting.estimated_time, '8:00 AM')

    def test_window_mode_shows_students_in_next_15_minutes(self):
        user = User.objects.create_user(username='screen-admin', password='12345678', is_staff=True)
        ScreenSettings.objects.create(
            room_count=1,
            waiting_count=5,
            estimate_time_per_student=5,
            exam_start_time=time(8, 0),
            public_screen_mode='window',
        )
        students = [
            Student.objects.create(name=f'Student {idx}', room='1', status='waiting', position=idx)
            for idx in range(6)
        ]
        captured = {}

        def capture_render(request, template_name, context):
            captured.update(context)
            return HttpResponse('ok')

        today = datetime.today().date()
        fake_now = make_aware(datetime.combine(today, time(8, 10)), timezone=get_current_timezone())

        def fake_localtime(value=None, timezone=None):
            if value is None:
                return fake_now
            return value

        self.client.force_login(user)
        with patch('screen.views.render', side_effect=capture_render), patch('screen.views.localtime', side_effect=fake_localtime):
            response = self.client.get(reverse('public_screen'))

        visible_numbers = [student.number for student in captured['schedule_students']]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(captured['screen_mode'], 'window')
        self.assertEqual(visible_numbers, [student.number for student in students[:4]])
        self.assertEqual(captured['schedule_window_start'], '08:00')
        self.assertEqual(captured['schedule_window_end'], '08:15')

    def test_window_mode_shifts_queue_and_shows_recent_finished_results(self):
        user = User.objects.create_user(username='screen-window-results-admin', password='12345678', is_staff=True)
        ScreenSettings.objects.create(
            room_count=1,
            waiting_count=5,
            estimate_time_per_student=5,
            exam_start_time=time(7, 30),
            result_display_seconds=30,
            public_screen_mode='window',
        )
        Student.objects.create(name='Finished Student', room='1', status='finished', position=0)
        active_students = [
            Student.objects.create(name=f'Active Student {idx}', room='1', status='waiting', position=idx + 1)
            for idx in range(4)
        ]
        finished_student = Student.objects.get(name='Finished Student')
        ExamResult.objects.create(
            number=finished_student.number,
            name=finished_student.name,
            grade=100,
            result='ناجح',
            room=1,
            sub_room='0',
        )
        captured = {}

        def capture_render(request, template_name, context):
            captured.update(context)
            return HttpResponse('ok')

        self.client.force_login(user)
        with patch('screen.views.render', side_effect=capture_render):
            response = self.client.get(reverse('public_screen'))

        active_numbers = [student.number for student in captured['schedule_students'] if not getattr(student, 'is_result', False)]
        result_rows = [student for student in captured['schedule_students'] if getattr(student, 'is_result', False)]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(active_numbers, [student.number for student in active_students])
        self.assertEqual(captured['schedule_window_start'], '07:30')
        self.assertEqual(captured['schedule_window_end'], '07:45')
        self.assertEqual(len(result_rows), 1)
        self.assertEqual(result_rows[0].display_status, '100 - ناجح')


class StudentImportMappingTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='admin-import', password='12345678', is_staff=True)
        self.client.force_login(self.user)
        ScreenSettings.objects.create(room_count=4, waiting_count=5)

    def make_upload(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['number', 'name', 'father', 'birth', 'institute', 'exam_type', 'parts'])
        for row in rows:
            sheet.append(row)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return SimpleUploadedFile(
            'students.xlsx',
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_student_form_allows_empty_room_for_random_assignment(self):
        form = StudentForm(data={
            'name': 'Random Room Student',
            'father_name': '',
            'birth_year': '',
            'institute_name': '',
            'exam_type': 'gh',
            'memorized_parts': '1',
            'room': '',
        })

        self.assertTrue(form.is_valid(), form.errors)
        self.assertIsNone(form.cleaned_data['room'])

    def test_upload_excel_leaves_students_in_assignment_pool(self):
        exam_type_label = next(iter(EXAM_TYPE_MAP.keys()))
        upload = self.make_upload([
            [None, 'Juz One', '', 2010, 'Institute', exam_type_label, 1],
            [None, 'Juz Two', '', 2010, 'Institute', exam_type_label, 2],
            [None, 'Juz Thirty A', '', 2010, 'Institute', exam_type_label, 30],
            [None, 'Juz Thirty B', '', 2010, 'Institute', exam_type_label, 30],
        ])

        response = self.client.post(reverse('upload_excel'), {
            'file': upload,
        })

        self.assertEqual(response.status_code, 302)
        students_by_name = {
            student.name: student
            for student in Student.objects.filter(name__startswith='Juz')
        }

        self.assertEqual({student.room for student in students_by_name.values()}, {'unassigned'})
        self.assertEqual({student.status for student in students_by_name.values()}, {'on_waiting_list'})

    def test_upload_excel_keeps_room_preferences_for_later_assignment(self):
        exam_type_label = next(iter(EXAM_TYPE_MAP.keys()))
        upload = self.make_upload([
            [None, 'Preferred Student', '', 2010, 'Institute', exam_type_label, 30],
        ])

        response = self.client.post(reverse('upload_excel'), {
            'file': upload,
            'mapping_juz_0': ['30'],
            'mapping_room_0': ['2', '4'],
        })
        student = Student.objects.get(name='Preferred Student')

        self.assertEqual(response.status_code, 302)
        self.assertEqual(student.room, 'unassigned')
        self.assertEqual(json.loads(student.preferred_rooms), [2, 4])

    def test_assign_imported_student_moves_one_student_to_a_room(self):
        student = Student.objects.create(
            name='Pool Student',
            room='unassigned',
            status='on_waiting_list',
            exam_type='gh',
        )

        response = self.client.post(reverse('assign_imported_student', kwargs={'student_number': student.number}))
        student.refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertIn(str(student.room), {'1', '2', '3', '4'})
        self.assertEqual(student.status, 'waiting')

    def test_assign_imported_student_balances_within_preferred_rooms(self):
        for index in range(6):
            Student.objects.create(name=f'Room 1 Student {index}', room='1', status='waiting', position=index)
        for index in range(2):
            Student.objects.create(name=f'Room 2 Student {index}', room='2', status='waiting', position=index)
        student = Student.objects.create(
            name='Preferred Pool Student',
            room='unassigned',
            status='on_waiting_list',
            exam_type='gh',
            preferred_rooms=json.dumps([1, 2]),
        )

        response = self.client.post(reverse('assign_imported_student', kwargs={'student_number': student.number}))
        student.refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertEqual(student.room, '2')
        self.assertEqual(student.preferred_rooms, '')

    def test_assign_all_imported_students_moves_entire_assignment_pool(self):
        students = [
            Student.objects.create(name=f'Bulk Pool Student {index}', room='unassigned', status='on_waiting_list')
            for index in range(3)
        ]
        Student.objects.create(name='Already Assigned', room='1', status='waiting')

        response = self.client.post(reverse('assign_all_imported_students'))
        for student in students:
            student.refresh_from_db()

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Student.objects.filter(room='unassigned').exists())
        self.assertEqual({student.status for student in students}, {'waiting'})
        self.assertTrue(all(parse_room_number(student.room) in {1, 2, 3, 4} for student in students))

    def test_remove_all_imported_students_deletes_only_assignment_pool(self):
        Student.objects.create(name='Pool Remove One', room='unassigned', status='on_waiting_list')
        Student.objects.create(name='Pool Remove Two', room='unassigned', status='on_waiting_list')
        assigned = Student.objects.create(name='Keep Assigned', room='1', status='waiting')

        response = self.client.post(reverse('remove_all_imported_students'))

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Student.objects.filter(room='unassigned').exists())
        self.assertTrue(Student.objects.filter(number=assigned.number).exists())
