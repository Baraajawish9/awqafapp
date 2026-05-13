import os
import csv
import json
import random
import openpyxl
from datetime import datetime, time, timedelta, date
from collections import Counter, defaultdict
import re
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.utils.encoding import smart_str
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.views.decorators.http import require_POST
from django.conf import settings
from django.db.models import Max
from django.utils.timezone import make_aware, localtime, get_current_timezone, now as timezone_now

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from django.db.models import Max  # <-- add this at the top of your file

from login.room_users import ensure_room_users
from .models import Student, ScreenSettings, ExamResult, STATUS_CHOICES, RoomQueue
from .forms import StudentForm, ScreenSettingsForm


STATUS_PRIORITY = {
    'in_exam': 0,
    'waiting': 1,
    'late': 2,
    'on_waiting_list': 3,
    'finished': 4,
}

ROOMS_PER_SCREEN_PAGE = 12
VISIBLE_STUDENTS_PER_TV_ROOM = 7
STUDENT_SLICE_ROTATE_MS = 5000


def contact(request):
    return render(request, 'screen/contact.html')


def get_tv_layout(room_count):
    if room_count <= 6:
        return {
            'rooms_per_page': 6,
            'visible_students': 9,
            'student_rotate_ms': 4500,
            'density_class': 'tv-density-6 tv-grid-3x2',
        }

    if room_count <= 8:
        return {
            'rooms_per_page': 8,
            'visible_students': 8,
            'student_rotate_ms': 4500,
            'density_class': 'tv-density-8 tv-grid-4x2',
        }

    if room_count <= 12:
        return {
            'rooms_per_page': 12,
            'visible_students': 7,
            'student_rotate_ms': 4500,
            'density_class': 'tv-density-12 tv-grid-4x3',
        }

    if room_count <= 15:
        return {
            'rooms_per_page': 15,
            'visible_students': 6,
            'student_rotate_ms': 4200,
            'density_class': 'tv-density-15 tv-grid-5x3',
        }

    if room_count == 16:
        return {
            'rooms_per_page': 16,
            'visible_students': 5,
            'student_rotate_ms': 4000,
            'density_class': 'tv-density-16 tv-grid-4x4',
        }

    if room_count <= 20:
        return {
            'rooms_per_page': 20,
            'visible_students': 5,
            'student_rotate_ms': 4000,
            'density_class': 'tv-density-20 tv-grid-5x4',
        }

    if room_count <= 24:
        return {
            'rooms_per_page': 24,
            'visible_students': 4,
            'student_rotate_ms': 3800,
            'density_class': 'tv-density-24 tv-grid-6x4',
        }

    return {
        'rooms_per_page': 30,
        'visible_students': 3,
        'student_rotate_ms': 3500,
        'density_class': 'tv-density-30 tv-grid-6x5',
    }


def parse_room_number(room_value):
    if room_value is None:
        return None

    room_text = str(room_value).strip().lower().replace('room', '')
    try:
        return int(room_text)
    except (TypeError, ValueError):
        return None


def room_sort_key(student):
    return (
        STATUS_PRIORITY.get(student.status, 99),
        student.position or 0,
        student.number or 0,
    )


def get_current_student_for_room(room_number):
    room_number = parse_room_number(room_number)
    if room_number is None:
        return None

    return (
        Student.objects
        .filter(room__in=[str(room_number), f'room{room_number}'], status='in_exam')
        .order_by('position', 'number')
        .first()
    )


def get_top_active_student_for_room(room_number):
    room_number = parse_room_number(room_number)
    if room_number is None:
        return None

    return (
        Student.objects
        .filter(room__in=[str(room_number), f'room{room_number}'])
        .exclude(status='finished')
        .order_by('position', 'number')
        .first()
    )


def is_current_student_for_room(student):
    room_number = parse_room_number(student.room)
    current_student = get_current_student_for_room(room_number)
    return bool(current_student and current_student.number == student.number)


def chunked(items, chunk_size):
    return [items[index:index + chunk_size] for index in range(0, len(items), chunk_size)]


def normalize_time_value(value, fallback=None):
    if isinstance(value, time):
        return value
    if isinstance(value, str):
        for time_format in ("%H:%M:%S", "%H:%M"):
            try:
                return datetime.strptime(value, time_format).time()
            except ValueError:
                pass
    return fallback or time(7, 0)


def get_room_count():
    settings = ScreenSettings.objects.last()
    return settings.room_count if settings else 5

def get_least_loaded_room():
    room_count = ScreenSettings.get_room_count()
    rooms = list(range(1, room_count + 1))
    random.shuffle(rooms)  # ✅ shuffle room order randomly

    # Count students in each room
    room_counts = Counter(Student.objects.values_list('room', flat=True))
    for r in rooms:
        room_counts.setdefault(r, 0)

    # Return the least loaded room (ties broken randomly)
    return min(rooms, key=lambda r: room_counts[r])


def get_least_loaded_room_from(room_options, extra_counts=None):
    rooms = [int(room) for room in room_options if room]
    if not rooms:
        return get_least_loaded_room()

    random.shuffle(rooms)
    room_counts = Counter(parse_room_number(room) for room in Student.objects.values_list('room', flat=True))
    if extra_counts:
        room_counts.update(extra_counts)

    for room in rooms:
        room_counts.setdefault(room, 0)

    return min(rooms, key=lambda room: room_counts[room])


def parse_import_room_mappings(post_data):
    mappings = []
    mapping_indexes = set()

    for key in post_data.keys():
        match = re.fullmatch(r'mapping_juz_(\d+)', key)
        if match:
            mapping_indexes.add(match.group(1))

    for index in sorted(mapping_indexes, key=int):
        juz_values = []
        room_values = []

        for value in post_data.getlist(f'mapping_juz_{index}'):
            try:
                juz_values.append(int(value))
            except (TypeError, ValueError):
                continue

        for value in post_data.getlist(f'mapping_room_{index}'):
            try:
                room_values.append(int(value))
            except (TypeError, ValueError):
                continue

        if juz_values and room_values:
            mappings.append({
                'juz': set(juz_values),
                'rooms': list(dict.fromkeys(room_values)),
            })

    return mappings


def get_mapped_import_room(parts, mappings, extra_counts):
    try:
        parts_number = int(parts)
    except (TypeError, ValueError):
        return None

    for mapping in mappings:
        if parts_number in mapping['juz']:
            room = get_least_loaded_room_from(mapping['rooms'], extra_counts)
            extra_counts[room] += 1
            return room

    return None


def get_room_loads(extra_counts=None):
    room_counts = Counter()
    for room in Student.objects.exclude(status='finished').values_list('room', flat=True):
        room_number = parse_room_number(room)
        if room_number is not None:
            room_counts[room_number] += 1

    if extra_counts:
        room_counts.update(extra_counts)

    return room_counts


def get_least_loaded_room_from(room_options, extra_counts=None):
    room_count = ScreenSettings.get_room_count()
    rooms = []

    for room in room_options:
        room_number = parse_room_number(room)
        if room_number is not None and 1 <= room_number <= room_count:
            rooms.append(room_number)

    rooms = list(dict.fromkeys(rooms))
    if not rooms:
        rooms = list(range(1, room_count + 1))

    random.shuffle(rooms)
    room_counts = get_room_loads(extra_counts)

    for room in rooms:
        room_counts.setdefault(room, 0)

    return min(rooms, key=lambda room: room_counts[room])


def get_least_loaded_room():
    return get_least_loaded_room_from(range(1, ScreenSettings.get_room_count() + 1))


def get_import_preferred_rooms(parts, mappings):
    try:
        parts_number = int(parts)
    except (TypeError, ValueError):
        return []

    for mapping in mappings:
        if parts_number in mapping['juz']:
            return mapping['rooms']

    return []


def get_student_preferred_rooms(student):
    try:
        rooms = json.loads(student.preferred_rooms or '[]')
    except (TypeError, ValueError, json.JSONDecodeError):
        return []

    if not isinstance(rooms, list):
        return []

    return rooms


@login_required
@staff_member_required
def public_screen(request):
    screen_settings = ScreenSettings.get_settings()
    room_count = screen_settings.room_count or 5
    result_display_seconds = screen_settings.result_display_seconds or 30
    tv_layout = get_tv_layout(room_count)
    rooms = list(range(1, room_count + 1))
    students = list(Student.objects.all().order_by('room', 'position', 'number'))
    students_by_number = {student.number: student for student in students}

    latest_ids = (
        ExamResult.objects
        .filter(sub_room='0')
        .values('number')
        .annotate(latest_id=Max('id'))
        .values_list('latest_id', flat=True)
    )
    results = ExamResult.objects.filter(id__in=latest_ids)
    latest_results = {
        res.number: {
            'grade': res.grade,
            'result': res.result
        }
        for res in results
    }

    for student in students:
        result = latest_results.get(student.number)
        student.latest_grade = result['grade'] if result else None
        student.latest_result = result['result'] if result else None
        student.latest_result_class = 'success' if student.latest_result == 'ناجح' else 'retry'
        student.room_number = parse_room_number(student.room)

    students_by_room = defaultdict(list)
    for student in students:
        if student.room_number in rooms and student.status != 'finished':
            students_by_room[student.room_number].append(student)

    for room_students in students_by_room.values():
        room_students.sort(key=room_sort_key)

    estimate_minutes = screen_settings.estimate_time_per_student or 5
    exam_start_time_value = normalize_time_value(screen_settings.exam_start_time, time(7, 0))
    naive_start_datetime = datetime.combine(datetime.today(), exam_start_time_value)
    tz = get_current_timezone()
    exam_start_time = localtime(make_aware(naive_start_datetime, timezone=tz))

    result_cutoff = timezone_now() - timedelta(seconds=result_display_seconds)
    recent_finished_by_room = defaultdict(list)
    seen_finished_numbers = set()
    for result in ExamResult.objects.filter(sub_room='0', room__in=rooms, timestamp__gte=result_cutoff).order_by('-timestamp', '-id'):
        student = students_by_number.get(result.number)
        if not student or student.status != 'finished' or result.number in seen_finished_numbers:
            continue
        seen_finished_numbers.add(result.number)
        result.room_number = parse_room_number(student.room) or result.room
        if result.room_number not in rooms:
            continue
        result.status = 'finished'
        result.latest_grade = result.grade
        result.latest_result = result.result
        result.result_class = 'success' if result.result == 'ناجح' else 'retry'
        result.latest_result_class = result.result_class
        result.estimated_time = None
        result.scheduled_time = 'نتيجة'
        result.display_status = f"{result.grade:g} - {result.result}"
        result.is_result = True
        recent_finished_by_room[result.room_number].append(result)

    if screen_settings.public_screen_mode == 'window':
        window_start = exam_start_time
        window_end = window_start + timedelta(minutes=15)
        schedule_students = []

        for room, room_students in students_by_room.items():
            active_students = sorted(
                [student for student in room_students if student.status != 'finished'],
                key=lambda student: (student.position or 0, student.number or 0),
            )
            for idx, student in enumerate(active_students):
                scheduled_time = exam_start_time + timedelta(minutes=idx * estimate_minutes)
                if window_start <= scheduled_time <= window_end:
                    student.scheduled_at = scheduled_time
                    student.scheduled_time = scheduled_time.strftime("%H:%M")
                    student.room_number = room
                    schedule_students.append(student)

        schedule_students.sort(key=lambda student: (
            getattr(student, 'scheduled_at', exam_start_time),
            student.room_number or 0,
            student.position or 0,
            student.number or 0,
        ))
        for room in rooms:
            schedule_students.extend(recent_finished_by_room.get(room, []))
        schedule_students_by_room = defaultdict(list)
        for student in schedule_students:
            schedule_students_by_room[student.room_number].append(student)

        schedule_room_cards = [
            {
                'number': room,
                'students': room_students,
                'total_count': len(room_students),
            }
            for room, room_students in sorted(schedule_students_by_room.items())
        ]
        schedule_room_pages = [
            {
                'number': index + 1,
                'rooms': page_rooms,
            }
            for index, page_rooms in enumerate(chunked(schedule_room_cards, tv_layout['rooms_per_page']))
        ]

        return render(request, 'screen/public_screen.html', {
            'screen_mode': 'window',
            'schedule_students': schedule_students,
            'schedule_room_pages': schedule_room_pages,
            'schedule_room_page_count': len(schedule_room_pages),
            'schedule_window_start': window_start.strftime("%H:%M"),
            'schedule_window_end': window_end.strftime("%H:%M"),
            'schedule_refresh_ms': 30000,
            'tv_density_class': tv_layout['density_class'],
        })

    for students_in_room in students_by_room.values():
        for idx, student in enumerate(students_in_room):
            scheduled_time = exam_start_time + timedelta(minutes=idx * estimate_minutes)
            student.estimated_time = scheduled_time.strftime("%I:%M %p").lstrip("0")

    result_cutoff = timezone_now() - timedelta(seconds=result_display_seconds)
    recent_finished_by_room = defaultdict(list)
    seen_finished_numbers = set()
    for result in ExamResult.objects.filter(sub_room='0', room__in=rooms, timestamp__gte=result_cutoff).order_by('-timestamp', '-id'):
        student = students_by_number.get(result.number)
        if not student or student.status != 'finished' or result.number in seen_finished_numbers:
            continue
        seen_finished_numbers.add(result.number)
        result.room_number = parse_room_number(student.room) or result.room
        if result.room_number not in rooms:
            continue
        result.status = 'finished'
        result.latest_grade = result.grade
        result.latest_result = result.result
        result.result_class = 'success' if result.result == 'ناجح' else 'retry'
        result.latest_result_class = result.result_class
        result.estimated_time = None
        recent_finished_by_room[result.room_number].append(result)

    room_cards = []
    for room in rooms:
        room_students = students_by_room.get(room, [])
        result_students = recent_finished_by_room.get(room, [])
        display_students = room_students + result_students
        status_counts = Counter(student.status for student in room_students)

        room_cards.append({
            'number': room,
            'students': display_students,
            'total_count': len(display_students),
            'active_count': len(room_students),
            'result_count': len(result_students),
            'current_student': next((s for s in room_students if s.status == 'in_exam'), None),
            'waiting_count': status_counts.get('waiting', 0) + status_counts.get('on_waiting_list', 0),
            'late_count': status_counts.get('late', 0),
        })

    room_pages = [
        {
            'number': index + 1,
            'rooms': page_rooms,
            'is_results': False,
        }
        for index, page_rooms in enumerate(chunked(room_cards, tv_layout['rooms_per_page']))
    ]

    return render(request, 'screen/public_screen.html', {
        'screen_mode': 'rooms',
        'room_pages': room_pages,
        'room_page_count': len(room_pages),
        'rotate_interval_ms': 1200,
        'student_slice_rotate_ms': tv_layout['student_rotate_ms'],
        'visible_students_per_room': tv_layout['visible_students'],
        'tv_density_class': tv_layout['density_class'],
        'exam_start_minutes': exam_start_time.hour * 60 + exam_start_time.minute,
        'estimate_minutes': estimate_minutes,
    })

@login_required
def clear_students(request):
    if request.method == 'POST':
        Student.objects.all().delete()
    return redirect('add_student')

def apply_automatic_status_for_room(room):
    settings = ScreenSettings.objects.last()
    waiting_limit = settings.waiting_count if settings else 5
    if parse_room_number(room) is None:
        return

    students = list(Student.objects.filter(room=room).exclude(status='finished').order_by('position', 'number'))
    late_students = [student for student in students if student.status == 'late']
    current_student = next((student for student in students if student.status == 'in_exam'), None)
    other_students = [
        student
        for student in students
        if student.status != 'late' and student != current_student
    ]
    new_order = []

    if current_student:
        new_order.append(current_student)

    waiting_students = other_students[:waiting_limit]
    new_order.extend(waiting_students)
    other_students = other_students[waiting_limit:]
    new_order.extend(late_students)
    new_order.extend(other_students)

    for idx, student in enumerate(new_order):
        student.position = idx
        if student in late_students:
            student.status = 'late'
        elif current_student and student == current_student:
            student.status = 'in_exam'
        elif student in waiting_students:
            student.status = 'waiting'
        else:
            student.status = 'on_waiting_list'
        student.save()


def update_student_status(request, student_number):
    if request.method == 'POST':
        student = get_object_or_404(Student, number=student_number)
        new_status = request.POST.get('status')
        old_room = student.room

        if new_status == "remove":
            messages.success(request, f"تم حذف الطالب {student.name}")
            student.delete()
            if parse_room_number(old_room) is not None:
                apply_automatic_status_for_room(old_room)
            return redirect('add_student')

        elif new_status.startswith("move:"):
            try:
                new_room = int(new_status.split(":")[1])
                student.room = new_room
                # Assign next available position in new room
                max_position = (
                    Student.objects.filter(room=new_room)
                    .aggregate(max_pos=Max('position'))['max_pos']
                )
                student.position = (max_position or 0) + 1
                student.save()
                messages.success(request, f"تم نقل الطالب {student.name} إلى لجنة {new_room}")
            except:
                messages.error(request, "حدث خطأ أثناء نقل الطالب")

        else:
            student.status = new_status
            student.save()
            messages.success(request, f"تم تحديث حالة الطالب {student.name} إلى {student.get_status_display()}")

        if parse_room_number(old_room) is not None:
            apply_automatic_status_for_room(old_room)
        if student.pk and student.room != old_room and parse_room_number(student.room) is not None:
            apply_automatic_status_for_room(student.room)
    
    return redirect('add_student')


@require_POST
@login_required
@staff_member_required
def assign_imported_student(request, student_number):
    student = get_object_or_404(Student, number=student_number)
    if parse_room_number(student.room) is not None:
        return redirect('add_student')

    assigned_room = assign_student_to_room(student)
    apply_automatic_status_for_room(student.room)
    messages.success(request, f"تم إدخال {student.name} إلى لجنة {assigned_room}")
    return redirect('add_student')


def assign_student_to_room(student, extra_counts=None):
    preferred_rooms = get_student_preferred_rooms(student)
    assigned_room = get_least_loaded_room_from(preferred_rooms, extra_counts)
    max_position = (
        Student.objects.filter(room__in=[str(assigned_room), f'room{assigned_room}'])
        .aggregate(max_pos=Max('position'))['max_pos']
    )
    student.room = str(assigned_room)
    student.position = (max_position or 0) + 1
    student.status = 'waiting'
    student.preferred_rooms = ''
    student.save()
    if extra_counts is not None:
        extra_counts[assigned_room] += 1
    return assigned_room


@require_POST
@login_required
@staff_member_required
def assign_all_imported_students(request):
    students = list(Student.objects.filter(room='unassigned').exclude(status='finished').order_by('number'))

    for student in students:
        assign_student_to_room(student)

    apply_automatic_status()
    messages.success(request, f"تم توزيع {len(students)} طالب.")
    return redirect('add_student')


@require_POST
@login_required
@staff_member_required
def remove_all_imported_students(request):
    deleted_count, _ = Student.objects.filter(room='unassigned').exclude(status='finished').delete()
    messages.success(request, f"تم حذف {deleted_count} طالب من قائمة انتظار التوزيع.")
    return redirect('add_student')



def apply_automatic_status():
    rooms = Student.objects.values_list('room', flat=True).distinct()

    for room in rooms:
        apply_automatic_status_for_room(room)


@login_required
def trigger_automatic_status(request):
    apply_automatic_status()
    return redirect('add_student')
@login_required
@staff_member_required
def add_student(request):
    form = StudentForm()

    if request.method == 'POST' and 'name' in request.POST:
        form = StudentForm(request.POST)
        if form.is_valid():
            student = form.save(commit=False)

            # Automatically assign the least loaded room only if not selected
            if not student.room:
                student.room = get_least_loaded_room()

            # Set position to the next available one in the assigned room
            max_position = (
                Student.objects.filter(room=student.room)
                .aggregate(max_pos=Max('position'))
            )['max_pos']
            student.position = (max_position or 0) + 1

            student.save()
            apply_automatic_status()

            messages.success(request, 'تمت إضافة الطالب بنجاح.')
            return redirect('add_student')
        else:
            print(form.errors)
            messages.error(request, 'الرجاء التحقق من صحة البيانات.')

    rooms = list(range(1, get_room_count() + 1))
    juz_list = list(range(1, 31))  # Juz numbers from 1 to 30

    search_query = request.GET.get('q', '').strip()
    selected_room_raw = request.GET.get('room', '').strip()
    selected_status = request.GET.get('status', '').strip()
    selected_room = parse_room_number(selected_room_raw)
    valid_statuses = {value for value, _label in STATUS_CHOICES}

    if selected_room not in rooms:
        selected_room = None

    if selected_status not in valid_statuses:
        selected_status = ''

    all_students = list(Student.objects.all().order_by('room', 'position', 'number'))
    students_by_room = {room: [] for room in rooms}
    unassigned_students = []

    for student in all_students:
        student.room_number = parse_room_number(student.room)
        if student.room_number in students_by_room:
            students_by_room[student.room_number].append(student)
        elif student.status != 'finished':
            student.preferred_room_numbers = get_student_preferred_rooms(student)
            unassigned_students.append(student)

    for room_students in students_by_room.values():
        room_students.sort(key=room_sort_key)

    def student_matches_filters(student):
        if selected_status and student.status != selected_status:
            return False

        if search_query:
            query = search_query.lower()
            searchable_values = [
                str(student.number or ''),
                student.name or '',
                student.father_name or '',
                student.institute_name or '',
            ]
            if not any(query in value.lower() for value in searchable_values):
                return False

        return True

    has_filters = bool(search_query or selected_room or selected_status)
    room_summaries = []
    room_cards = []

    for room in rooms:
        room_students = students_by_room.get(room, [])
        status_counts = Counter(student.status for student in room_students)
        matching_students = [
            student
            for student in room_students
            if (selected_room is None or room == selected_room) and student_matches_filters(student)
        ]
        current_student = next((student for student in room_students if student.status == 'in_exam'), None)

        room_summaries.append({
            'number': room,
            'total_count': len(room_students),
            'match_count': len(matching_students),
            'current_student': current_student,
            'waiting_count': status_counts.get('waiting', 0) + status_counts.get('on_waiting_list', 0),
            'late_count': status_counts.get('late', 0),
        })

        should_show_room = (
            (not has_filters and room_students)
            or (not has_filters and room == rooms[0] and not all_students)
            or (has_filters and bool(matching_students))
            or (has_filters and selected_room == room)
        )

        if should_show_room:
            room_cards.append({
                'number': room,
                'students': matching_students if has_filters else room_students,
                'total_count': len(room_students),
                'match_count': len(matching_students),
            })

    if not room_cards and rooms and not has_filters:
        room_cards.append({
            'number': rooms[0],
            'students': [],
            'total_count': 0,
            'match_count': 0,
        })

    all_institutes = (
        Student.objects.values_list('institute_name', flat=True)
        .distinct()
        .order_by('institute_name')
    )

    context = {
        'form': form,
        'rooms': rooms,
        'juz_list': juz_list,
        'students_by_room': students_by_room,
        'unassigned_students': unassigned_students,
        'room_summaries': room_summaries,
        'room_cards': room_cards,
        'search_query': search_query,
        'selected_room': selected_room or '',
        'selected_status': selected_status,
        'has_filters': has_filters,
        'status_choices': STATUS_CHOICES,
        'all_institutes': all_institutes,
    }

    return render(request, 'screen/add_student.html', context)



@require_POST
@login_required
def submit_grade(request, student_number, subroom):
    student = get_object_or_404(Student, number=student_number)
    if subroom not in (1, 2):
        return HttpResponse("Invalid subroom number", status=400)

    room_number = parse_room_number(student.room)
    if room_number is None:
        return HttpResponse("Invalid room", status=400)

    if not is_current_student_for_room(student):
        return redirect('room_view', room_name=f'room{room_number}', subroom=subroom)
    
    # Parse exam config
    if student.exam_type == "gh":
        pass_threshold = 80
    elif student.exam_type == "nz":
        pass_threshold = 90
    else:
        pass_threshold = 0

    # Get submitted final grade for this subroom from form
    try:
        final_grade = float(request.POST.get("final_grade", "100"))
    except ValueError:
        final_grade = 100
    final_grade = max(0, min(100, final_grade))

    # Save/update this subroom's partial result with a temporary 'result' value
    exam_result, created = ExamResult.objects.update_or_create(
        number=student.number,
        sub_room=str(subroom),
        defaults={
            'name': student.name,
            'grade': final_grade,
            'result': 'قيد التقييم',  # temporary placeholder to satisfy NOT NULL
            'room': room_number,
        }
    )

    # Check both branch grades for this student.
    subroom_results = ExamResult.objects.filter(number=student.number, sub_room__in=['1', '2'])
    grades_by_subroom = {str(er.sub_room): er.grade for er in subroom_results}

    if len(grades_by_subroom) == 2:
        avg_grade = sum(grades_by_subroom.values()) / 2
        result = "ناجح" if avg_grade >= pass_threshold else "إعادة"

        # Save final summary with sub_room=0 to indicate final average
        ExamResult.objects.update_or_create(
            number=student.number,
            sub_room='0',
            defaults={
                'name': student.name,
                'grade': avg_grade,
                'result': result,
                'room': room_number,
            }
        )

        # Mark student finished
        student.status = 'finished'
        student.save()
        for key in ['locked_student_number', 'locked_room', 'locked_subroom']:
            request.session.pop(key, None)

        # Log to CSV after final grade calculated
        csv_path = os.path.join(settings.BASE_DIR, 'grades_log.csv')
        file_exists = os.path.isfile(csv_path)
        with open(csv_path, mode='a', encoding='utf-8', newline='') as f:
            writer = csv.writer(f)
            if not file_exists:
                writer.writerow(['Student Number', 'Name', 'Final Grade', 'Result', 'Sub Room'])
            writer.writerow([student.number, student.name, avg_grade, result, 'final_average'])

        return redirect('room_view', room_name=f'room{room_number}', subroom=subroom)

    return redirect('room_view', room_name=f'room{room_number}', subroom=subroom)

@login_required
@staff_member_required
def edit_settings(request):
    old_settings = ScreenSettings.get_settings()
    old_room_count = old_settings.room_count if old_settings else 0

    if request.method == 'POST':
        form = ScreenSettingsForm(request.POST, instance=old_settings)
        if form.is_valid():
            new_settings = form.save(commit=False)
            new_room_count = new_settings.room_count

            if new_room_count != old_room_count:
                students = list(Student.objects.all().exclude(status="finished"))
                room_positions = {room: 0 for room in range(1, new_room_count + 1)}

                for student in students:
                    # Shuffle room list each time to add randomness
                    rooms = list(room_positions.keys())
                    random.shuffle(rooms)

                    # Find the room with the minimal position count among shuffled rooms
                    min_room = min(rooms, key=lambda r: room_positions[r])

                    room_positions[min_room] += 1
                    student.room = min_room
                    student.position = room_positions[min_room]
                    student.save()

            new_settings.save()
            ensure_room_users(new_settings.room_count)

            apply_automatic_status()

            return redirect('add_student')

    else:
        form = ScreenSettingsForm(instance=old_settings)

    return render(request, 'screen/edit_settings.html', {'form': form})



EXAM_TYPE_MAP = {
    "غيباً": "gh",
    "نظراً": "nz",
}

@login_required
@require_POST
def upload_excel(request):
    file = request.FILES.get("file")
    if not file or not file.name.endswith(".xlsx"):
        return redirect('add_student')

    room_mappings = parse_import_room_mappings(request.POST)
    wb = openpyxl.load_workbook(file)
    sheet = wb.active

    current_number = Student.objects.aggregate(Max('number'))['number__max'] or 0

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue

        name = str(row[1]).strip() if row[1] else ''
        if not name or "الاسم" in name:
            continue

        try:
            father_name = str(row[2]).strip() if row[2] else ''
            birth_year_raw = row[3]
            birth_year = None
            if birth_year_raw:
                try:
                    birth_year = int(birth_year_raw)
                except Exception as e:
                    print(f"Invalid birth year in row: {row} -> {e}")
                    birth_year = None

            institute = str(row[4]).strip() if row[4] else ''
            exam_type_arabic = str(row[5]).strip() if row[5] else ''
            exam_type = EXAM_TYPE_MAP.get(exam_type_arabic)
            parts = int(row[6]) if row[6] else 0

            if not exam_type:
                print(f"Unknown exam type: {exam_type_arabic} in row: {row}")
                continue

            current_number += 1

            assigned_room = 'unassigned'
            preferred_rooms = get_import_preferred_rooms(parts, room_mappings)

            Student.objects.create(
                number=current_number,
                name=name,
                father_name=father_name,
                birth_year=birth_year,
                institute_name=institute,
                exam_type=exam_type,
                memorized_parts=parts,
                room=assigned_room,
                status='on_waiting_list',
                preferred_rooms=json.dumps(preferred_rooms) if preferred_rooms else '',
            )

        except Exception as e:
            print("Error in row:", row, str(e))
            continue

    return redirect('add_student')



def remove_student(request, student_number):
    if request.method == 'POST':
        student = get_object_or_404(Student, number=student_number)
        student.delete()
        messages.success(request, "تم حذف الطالب بنجاح.")
    return redirect('add_student')  

@require_POST
@login_required
def clear_all_results(request):
    # Delete all ExamResult entries
    ExamResult.objects.all().delete()

    # Delete only students with status "finished"
    Student.objects.filter(status="finished").delete()

    # Redirect to the add-student page after clearing
    return redirect('add_student')

@login_required
def export_students_excel(request):
    institute_name = request.GET.get('institute')
    if not institute_name:
        return HttpResponse("يرجى اختيار اسم المعهد", status=400)
    include_grade = request.GET.get('include_grade') == '1'
    last_column = 10 if include_grade else 8

    settings = ScreenSettings.get_settings()
    estimated_time_per_student = settings.estimate_time_per_student or 5
    exam_start_time = normalize_time_value(settings.exam_start_time, time(8, 0))
    start_minutes = exam_start_time.hour * 60 + exam_start_time.minute

    # Build room-wise queues
    students = Student.objects.all().order_by('room', 'number')
    room_queues = defaultdict(list)
    for student in students:
        room_queues[student.room].append(student)

    final_results = {}
    if include_grade:
        latest_result_ids = (
            ExamResult.objects
            .filter(sub_room='0')
            .values('number')
            .annotate(latest_id=Max('id'))
            .values_list('latest_id', flat=True)
        )
        final_results = {
            result.number: result
            for result in ExamResult.objects.filter(id__in=latest_result_ids)
        }

    # Assign estimated starts based on position in room queue.
    student_start_minutes = {}
    if room_queues:
        max_queue_length = max(len(queue) for queue in room_queues.values())
        for index in range(max_queue_length):
            for room in sorted(room_queues.keys()):
                queue = room_queues[room]
                if index < len(queue):
                    student = queue[index]
                    est_minutes = start_minutes + index * estimated_time_per_student
                    student_start_minutes[student.number] = est_minutes

    def format_estimated_minutes(minutes):
        return f"{minutes // 60:02}:{minutes % 60:02}"

    def estimated_range_for_students(export_students):
        starts = [
            student_start_minutes[student.number]
            for student in export_students
            if student.number in student_start_minutes
        ]
        if not starts:
            return ''

        range_start = min(starts)
        range_end = max(starts) + estimated_time_per_student
        return f"{format_estimated_minutes(range_start)} - {format_estimated_minutes(range_end)}"

    # Create Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.sheet_view.rightToLeft = True
    current_row = 1

    def write_headers():
        nonlocal current_row
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=last_column)
        ws.cell(row=current_row, column=1, value="الجمهورية العربية السورية").alignment = Alignment(horizontal='center')
        ws.cell(row=current_row, column=1).font = Font(size=14, bold=True)
        current_row += 1

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=last_column)
        ws.cell(row=current_row, column=1, value="استمارة  اختبار الأجزاء المتفرقة").alignment = Alignment(horizontal='center')
        ws.cell(row=current_row, column=1).font = Font(size=12, bold=True)
        current_row += 1

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        ws.cell(row=current_row, column=1, value="وزارة الأوقاف").alignment = Alignment(horizontal='center')
        ws.cell(row=current_row, column=1).font = Font(bold=True)
        ws.merge_cells(start_row=current_row, start_column=5, end_row=current_row, end_column=last_column)
        ws.cell(row=current_row, column=5, value="مركز الحسنين").alignment = Alignment(horizontal='center')
        ws.cell(row=current_row, column=5).font = Font(bold=True)
        current_row += 1

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=last_column)
        ws.cell(row=current_row, column=1, value="مديرية معاهد تحفيظ القرآن الكريم بدمشق").alignment = Alignment(horizontal='center')
        ws.cell(row=current_row, column=1).font = Font(bold=True)
        current_row += 2

    def write_table_header():
        nonlocal current_row
        headers = ['الرقم', 'الاسم والكنية', 'اسم الأب', 'تاريخ الولادة', 'اسم المعهد', 'غيباً/نظراً', 'الأجزاء المحفوظة', 'الوقت المتوقع', 'الدرجة', 'النتيجة']
        if not include_grade:
            headers = headers[:8]
        for i, header in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=i, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
        current_row += 1

    def write_student_row(student, estimated_time):
        nonlocal current_row
        exam_type = student.get_exam_type_display() if hasattr(student, 'get_exam_type_display') else student.exam_type
        final_result = final_results.get(student.number)
        row = [
            student.number,
            smart_str(student.name),
            smart_str(student.father_name or ''),
            student.birth_year if student.birth_year else '',
            smart_str(student.institute_name or ''),
            exam_type,
            student.memorized_parts or '',
            estimated_time,
        ]
        if include_grade:
            row.extend([
                final_result.grade if final_result else '',
                smart_str(final_result.result) if final_result else '',
            ])
        ws.append(row)
        current_row += 1

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 15
    if include_grade:
        ws.column_dimensions['I'].width = 12
        ws.column_dimensions['J'].width = 12

    write_headers()
    write_table_header()

    if institute_name == '__all__':
        grouped = defaultdict(list)
        for s in students:
            grouped[s.institute_name].append(s)

        for institute in sorted(grouped.keys(), key=lambda value: value or ''):
            estimated_time = estimated_range_for_students(grouped[institute])
            sorted_students = sorted(grouped[institute], key=lambda x: student_start_minutes.get(x.number, 0))
            for stu in sorted_students:
                write_student_row(stu, estimated_time)
        filename_prefix = "All_Students_With_Grades" if include_grade else "All_Students"
        filename = f"{filename_prefix}_{datetime.today().date()}.xlsx"
    else:
        filtered_students = list(Student.objects.filter(institute_name=institute_name).order_by('room', 'number'))
        estimated_time = estimated_range_for_students(filtered_students)
        for student in sorted(filtered_students, key=lambda x: student_start_minutes.get(x.number, 0)):
            write_student_row(student, estimated_time)
        filename_prefix = "Students_With_Grades" if include_grade else "Students"
        filename = f"{filename_prefix}_{institute_name}_{datetime.today().date()}.xlsx"


    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

@require_POST
@login_required
def move_student_position(request, number):
    direction = request.POST.get("direction")
    student = get_object_or_404(Student, number=number)

    same_room_students = list(Student.objects.filter(room=student.room).order_by('position'))

    try:
        index = same_room_students.index(student)
    except ValueError:
        return redirect('add_student')

    if direction == "up" and index > 0:
        same_room_students[index], same_room_students[index - 1] = same_room_students[index - 1], same_room_students[index]
    elif direction == "down" and index < len(same_room_students) - 1:
        same_room_students[index], same_room_students[index + 1] = same_room_students[index + 1], same_room_students[index]

    for idx, s in enumerate(same_room_students):
        s.position = idx
        s.save()

    apply_automatic_status()  # reassign statuses based on new order

    return redirect('add_student')
